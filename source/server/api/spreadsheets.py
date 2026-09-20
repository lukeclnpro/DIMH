from pathlib import Path
import csv, io
from fastapi import APIRouter, HTTPException, Request
from server.security.auth import require_user
from pydantic import BaseModel
from server.services.file_service import safe_path

router=APIRouter(prefix='/api/spreadsheets')
class Sheet(BaseModel): path:str; rows:list[list[str]]

@router.get('')
def read(path:str):
    p=safe_path(path)
    if p.suffix.lower() not in {'.csv','.xlsx','.xls'} or not p.is_file(): raise HTTPException(400,'Tableur non pris en charge')
    if p.suffix.lower()!='.csv':
        try:
            from openpyxl import load_workbook
            wb=load_workbook(p, data_only=False); ws=wb.active; rows=[[str(c.value or '') for c in r] for r in ws.iter_rows()]; return {'path':path,'rows':rows}
        except Exception as e: raise HTTPException(400,f'Impossible de lire le tableur: {e}')
    with p.open('r',encoding='utf-8-sig',newline='') as f: return {'path':path,'rows':list(csv.reader(f))}

@router.put('')
def write(data:Sheet, request: Request):
    if require_user(request)['public']: raise HTTPException(403, 'Le compte public est en lecture seule. Connectez-vous pour modifier le Drive.')
    p=safe_path(data.path)
    if p.suffix.lower() not in {'.csv','.xlsx','.xls'}: raise HTTPException(400,'Extension non prise en charge')
    p.parent.mkdir(parents=True,exist_ok=True)
    if p.suffix.lower()!='.csv':
        try:
            from openpyxl import Workbook
            wb=Workbook(); ws=wb.active
            for row in data.rows: ws.append(row)
            wb.save(p); return {'ok':True}
        except Exception as e: raise HTTPException(400,f'Impossible d’écrire le tableur: {e}')
    with p.open('w',encoding='utf-8',newline='') as f: csv.writer(f).writerows(data.rows)
    return {'ok':True}
